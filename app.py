from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from datetime import date
from urllib.parse import urlparse
import sqlite3
import json
import math
import csv
import io
import os
from finance import replay

ROOT = Path(__file__).parent
DB = Path(os.environ.get('FINANCE_DB', ROOT / 'data' / 'finance.db'))

def connect():
    DB.parent.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    return c

def initialize():
    with connect() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS settings(id INTEGER PRIMARY KEY CHECK(id=1), initial REAL NOT NULL, start TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS assets(id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE);
        CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY, asset_id INTEGER NOT NULL REFERENCES assets(id),
          day TEXT NOT NULL, kind TEXT NOT NULL, amount REAL NOT NULL, price REAL NOT NULL, fee REAL NOT NULL, note TEXT NOT NULL);
        ''')

def state(c):
    setting = c.execute('SELECT * FROM settings').fetchone()
    assets = [dict(x) for x in c.execute('SELECT * FROM assets ORDER BY id')]
    events = [dict(x) for x in c.execute('SELECT * FROM events ORDER BY day,id')]
    if not setting:
        return dict(configured=False, assets=assets)
    return dict(configured=True, assets=assets, **replay(setting['initial'], setting['start'], assets, events))

def number(data, key, minimum=0, strict=False):
    v = float(data.get(key, 0))
    if not math.isfinite(v) or v > 1e12 or (v <= minimum if strict else v < minimum) or (key == 'price' and v < 1e-8):
        raise ValueError(f'{key} 数值超出有效范围')
    return v

def valid_day(value):
    d = date.fromisoformat(value)
    if not date(2000,1,1) <= d <= date.today():
        raise ValueError('日期须在 2000-01-01 至今天之间')
    return d.isoformat()

class Handler(BaseHTTPRequestHandler):
    def send(self, status, body, kind='application/json; charset=utf-8', filename=None):
        if not isinstance(body, bytes):
            body = json.dumps(body, ensure_ascii=False, allow_nan=False).encode()
        self.send_response(status)
        self.send_header('Content-Type', kind)
        self.send_header('Content-Length', str(len(body)))
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Cache-Control', 'no-store')
        if filename:
            self.send_header('Content-Disposition', f'attachment; filename="{filename}"')
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlparse(self.path).path
        if path in ['/', '/static/app.js', '/static/style.css']:
            file = ROOT / ('static/index.html' if path=='/' else path[1:])
            kind = {'html':'text/html','js':'text/javascript','css':'text/css'}[file.suffix[1:]]
            return self.send(200, file.read_bytes(), kind+'; charset=utf-8')
        if path not in ['/api/state','/api/export.csv','/api/export.xlsx']:
            return self.send(404, {'error':'页面不存在'})
        with connect() as c:
            s = state(c)
        if path=='/api/state':
            return self.send(200,s)
        if not s['configured']:
            return self.send(400, {'error':'请先设置投资本金'})
        rows = [['日期','标的','类型','成交金额','单价/净值','手续费','剩余份额','标的持仓市值','已实现收益','现金余额','总资产','总收益','总收益率(%)','备注']]
        names = dict(buy='买入',sell='卖出',quote='估值')
        for e in s['ledger']:
            rows.append([e['day'],e['name'],names[e['kind']],e['amount'],e['price'],e['fee'],e['units'],e['holding'],e['realized'],e['cash'],e['total'],e['profit'],e['return_pct'], e['note']])
        # Prevent text cells from becoming spreadsheet formulas.
        rows = [[("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v) for v in row] for row in rows]
        if path.endswith('.csv'):
            stream=io.StringIO(); csv.writer(stream).writerows(rows)
            return self.send(200,stream.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8','finance.csv')
        from openpyxl import Workbook
        book=Workbook(); sheet=book.active; sheet.title='投资流水'
        for row in rows: sheet.append(row)
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        daily=book.create_sheet('每日资产')
        daily.append(['日期','总资产','现金余额','累计收益']+[p['name']+'市值' for p in s['holdings']])
        for r in s['curve']: daily.append([r['day'],r['total'],r['cash'],r['profit']]+[r['values'][str(p['id'])] for p in s['holdings']])
        summary=book.create_sheet('持仓汇总'); summary.append(['标的','份额','最新净值','剩余成本','持仓市值','已实现收益','未实现收益','累计收益'])
        summary.append(['余额宝',None,None,None,s['cash'],s['interest'],0,s['interest']])
        for p in s['holdings']: summary.append([p[k] for k in ['name','units','price','cost','value','realized','unrealized','profit']])
        for ws in book.worksheets:
            for row in ws:
                for cell in row:
                    if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@','\t','\r')):
                        cell.value="'"+cell.value
        stream=io.BytesIO(); book.save(stream)
        self.send(200,stream.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','finance.xlsx')

    def do_POST(self):
        # Local personal tool: prevent cross-origin browser writes.
        origin=self.headers.get('Origin')
        if origin and urlparse(origin).netloc != self.headers.get('Host'):
            return self.send(403, {'error':'不允许跨站写入'})
        try:
            length=int(self.headers.get('Content-Length',0))
            if length > 16384: raise ValueError('请求过大')
            if self.headers.get_content_type() != 'application/json': raise ValueError('需要 JSON 请求')
            data=json.loads(self.rfile.read(length))
            if not isinstance(data,dict): raise ValueError('请求内容必须为对象')
            with connect() as c:
                path=urlparse(self.path).path
                if path=='/api/settings':
                    if c.execute('SELECT 1 FROM settings').fetchone(): raise ValueError('本金已初始化；为保留账本一致性，不支持直接修改')
                    c.execute('INSERT INTO settings VALUES(1,?,?)',(number(data,'initial',strict=True),valid_day(data['start'])))
                elif path=='/api/assets':
                    name=str(data['name']).strip()
                    if not name or len(name)>60 or name=='余额宝': raise ValueError('标的名称须为 1–60 字，且不能为余额宝')
                    c.execute('INSERT INTO assets(name) VALUES(?)',(name,))
                elif path=='/api/events':
                    setting=c.execute('SELECT * FROM settings').fetchone()
                    if not setting: raise ValueError('请先初始化本金')
                    day=valid_day(data['day'])
                    if day<setting['start']: raise ValueError('流水不能早于投资起始日期')
                    kind=data['kind']
                    if kind not in ['buy','sell','quote']: raise ValueError('无效的操作类型')
                    amount=number(data,'amount',strict=kind!='quote') if kind!='quote' else 0
                    fee=number(data,'fee') if kind!='quote' else 0
                    note=str(data.get('note','')).strip()
                    if len(note)>500: raise ValueError('备注最多 500 字')
                    c.execute('INSERT INTO events(asset_id,day,kind,amount,price,fee,note) VALUES(?,?,?,?,?,?,?)',
                              (int(data['asset_id']),day,kind,amount,number(data,'price',strict=True),fee,note))
                elif path=='/api/delete':
                    c.execute('DELETE FROM events WHERE id=?',(int(data['id']),))
                else:
                    return self.send(404,{'error':'接口不存在'})
                result=state(c)  # Revalidate the entire history; any failure rolls back.
            self.send(200,result)
        except (ValueError,KeyError,TypeError,sqlite3.IntegrityError) as e:
            self.send(400,{'error': '标的名称重复或标的不存在' if isinstance(e,sqlite3.IntegrityError) else str(e)})

if __name__=='__main__':
    initialize()
    server=ThreadingHTTPServer((os.environ.get('HOST','127.0.0.1'),int(os.environ.get('PORT','8000'))),Handler)
    print(f'Finance listening on {server.server_address}',flush=True)
    server.serve_forever()
