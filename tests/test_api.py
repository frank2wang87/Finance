import unittest
import tempfile
from pathlib import Path
from datetime import date
import threading
import urllib.request
import urllib.error
import json
import io
from openpyxl import load_workbook
import app

class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.previous=app.DB
        app.DB=Path(cls.temp.name)/'test.db'
        app.initialize()
        cls.server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True)
        cls.thread.start()
        cls.base=f'http://127.0.0.1:{cls.server.server_port}'

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();cls.thread.join()
        app.DB=cls.previous;cls.temp.cleanup()

    def request(self,path,data=None,headers=None):
        req=urllib.request.Request(self.base+path, data=json.dumps(data).encode() if data is not None else None,
            headers=headers or ({'Content-Type':'application/json'} if data is not None else {}))
        try:
            with urllib.request.urlopen(req) as r: return r.status,r.read()
        except urllib.error.HTTPError as e:
            return e.code,e.read()

    def post(self,path,data):
        code,body=self.request(path,data)
        self.assertEqual(code,200,body)
        return json.loads(body)

    def test_full_workflow_and_exports(self):
        today=date.today().isoformat()
        self.assertIn('投资手记',self.request('/')[1].decode())
        self.post('/api/settings',dict(initial=10000,start=today))
        s=self.post('/api/assets',dict(name='=1+1'))
        asset=s['assets'][0]['id']
        trade=dict(asset_id=asset,day=today,kind='buy',amount=4000,price=10,fee=10,note='=SUM(1,2)')
        s=self.post('/api/events',trade)
        buyid=s['ledger'][0]['id']
        s=self.post('/api/events',dict(trade,kind='quote',price=12))
        s=self.post('/api/events',dict(trade,kind='sell',amount=1200,price=12,fee=5))
        self.assertEqual(s['total'],10785)
        self.assertEqual(s['holdings'][0]['units'],300)
        self.assertEqual(s['holdings'][0]['realized'],192.5)
        status,_=self.request('/api/delete',dict(id=buyid))
        self.assertEqual(status,400)
        self.assertEqual(len(json.loads(self.request('/api/state')[1])['ledger']),3)
        status,_=self.request('/api/events',dict(trade,amount=999999))
        self.assertEqual(status,400)
        self.assertEqual(len(json.loads(self.request('/api/state')[1])['ledger']),3)
        status,csv=self.request('/api/export.csv')
        self.assertEqual(status,200);self.assertTrue(csv.startswith(b'\xef\xbb\xbf'))
        self.assertIn("'=SUM(1,2)",csv.decode('utf-8-sig'))
        status,excel=self.request('/api/export.xlsx')
        self.assertEqual(status,200)
        wb=load_workbook(io.BytesIO(excel))
        self.assertEqual(wb.sheetnames,['投资流水','每日资产','持仓汇总'])
        self.assertEqual(wb['每日资产'].cell(2,2).value,10785)
        self.assertTrue(all(c.data_type!='f' for ws in wb for row in ws for c in row))
        status,_=self.request('/api/assets',dict(name='bad'),{'Content-Type':'application/json','Origin':'https://evil.example'})
        self.assertEqual(status,403)
        s=self.post('/api/delete',dict(id=s['ledger'][-1]['id']))
        self.assertEqual(len(s['ledger']),2)
        self.assertEqual(s['total'],10790)
        status,_=self.request('/api/events',dict(trade,price=1e-300))
        self.assertEqual(status,400)

if __name__=='__main__': unittest.main()
